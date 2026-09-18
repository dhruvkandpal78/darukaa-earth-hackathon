"""Builds the Word submission overview; real URLs can be supplied after external deployment."""

import argparse
from pathlib import Path
from urllib.parse import urlparse
from docx import Document
from docx.shared import Inches, Pt, RGBColor


def create_submission(repository: str, demo: str, draft: bool = False):
    """Keep the review package concise and mark unavailable links rather than inventing them."""
    # A final package must never quietly replace required live links with placeholders.
    if not draft:
        for label, value in [("repository", repository), ("live demo", demo)]:
            parsed = urlparse(value)
            if (
                parsed.scheme != "https"
                or not parsed.netloc
                or parsed.hostname in {"localhost", "127.0.0.1"}
            ):
                raise ValueError(f"Supply a real public HTTPS {label} URL, or use --draft.")
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.75)
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2"]:
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Normal"].font.size = Pt(11)
    document.styles["Normal"].paragraph_format.space_after = Pt(7)
    document.styles["Title"].font.size = Pt(25)
    document.styles["Heading 1"].font.size = Pt(13)
    document.add_paragraph("Darukaa Earth project submission", "Title")
    document.add_paragraph("Full stack geospatial analytics platform", "Subtitle")
    document.add_paragraph(
        "Darukaa Earth organizes carbon and biodiversity projects, maps their geographic sites, and displays performance over time. This overview gives reviewers the application links, technical structure, setup path, and deployment requirements."
    )
    document.add_heading("Review links", 1)
    document.add_paragraph(f"Private GitHub repository: {repository}")
    document.add_paragraph(f"Public live demo: {demo}")
    document.add_paragraph(
        "API documentation after deployment: append /api/docs to the live demo URL. Local sample preview: http://127.0.0.1:5173. Full local application: http://127.0.0.1:8000."
    )
    document.add_heading("Architecture and schema", 1)
    document.add_paragraph(
        "React and TypeScript provide the interface. Mapbox GL JS displays and draws geographic boundaries when a public Mapbox token is configured; a MapLibre/OpenStreetMap preview fallback works without that token. Chart.js presents carbon, biodiversity, and canopy time series. Python FastAPI validates requests and enforces account ownership. PostgreSQL with PostGIS stores and measures polygons."
    )
    document.add_paragraph(
        "Four related tables hold users, projects, sites, and monthly observations. Projects belong to users; sites belong to projects; observations belong to sites. Site geometry uses EPSG 4326 with a spatial index. PostGIS calculates hectares from geographic area. A unique site/date constraint prevents duplicate observations. Passwords use Argon2 hashes; JWTs expire after one hour."
    )
    document.add_heading("Local setup", 1)
    document.add_paragraph(
        "Copy .env.example to .env and supply a random JWT_SECRET of at least 32 characters. With Docker Compose installed, run docker compose up --build, then open port 8000. For the sample interface alone, use Node.js 22 and run npm ci followed by npm run dev. Native backend development uses Python 3.12, backend/requirements.txt, a PostgreSQL/PostGIS database, and python -m backend.migrate before Uvicorn starts."
    )
    document.add_heading("Quality and deployment", 1)
    document.add_paragraph(
        "Husky and lint-staged format and lint staged changes. GitHub Actions checks Prettier, ESLint, Ruff, frontend tests, backend tests against a PostGIS service, and a production build. A successful main-branch run calls the Render deployment hook with the tested commit SHA. The Docker image serves the React build and Python API from one origin. render.yaml defines the web service and private database; the deploy-hook secret must be configured in GitHub."
    )
    document.add_heading("Reviewer access and data", 1)
    document.add_paragraph(
        "For a private repository, grant access to ankita.dasgupta@darukaa.com, harsh.kumar@darukaa.com, utkarsh.gauniyal@darukaa.com, and guneet.mutreja@darukaa.com. Reviewers can register their own application account; no shared production password is supplied. The sample portfolio uses explicitly labeled synthetic observations and approximate boundaries, not certified carbon or biodiversity measurements."
    )
    document.add_heading("Completion notes", 1)
    if draft:
        document.add_paragraph(
            "Repository access and live deployment remain pending account configuration. This draft must not be submitted until the real links are included and deployment is verified. The source includes deployment configuration and a requirement checklist in README.md."
        )
    else:
        document.add_paragraph(
            "Open the live demo above to explore the labeled sample portfolio. Register an account to create a private project, add a polygon site, and reopen its details. Newly created sites correctly show no observations until an operator imports measurements. The repository README contains the complete setup, data assumptions, and verification record."
        )
    document.add_heading("Submission instructions", 1)
    document.add_paragraph(
        "Submit this single Word document through My Jobs or the Applied Job page for the Darukaa Earth role, using its document submission option. Confirm that the repository and live demo links are accessible to the hiring team before uploading."
    )
    filename = "Darukaa_Earth_Submission_DRAFT.docx" if draft else "Darukaa_Earth_Submission.docx"
    destination = Path(__file__).resolve().parent.parent / "output" / filename
    destination.parent.mkdir(exist_ok=True)
    document.save(destination)
    return destination


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repository", default="PENDING - authenticated GitHub account required")
    parser.add_argument("--demo", default="PENDING - hosting deployment required")
    parser.add_argument(
        "--draft",
        action="store_true",
        help="Generate an explicitly named draft while deployment links are unavailable.",
    )
    args = parser.parse_args()
    print(create_submission(args.repository, args.demo, args.draft))
